import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import wbs_data as D  # noqa: E402

OUT = sys.argv[1]
FONT = "Arial"
NAVY = "1B2B45"
ORANGE = "F26B3A"
thin = Side(style="thin", color="CBD5E1")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
SUB = Font(name=FONT, size=10, italic=True, color="475569")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
FILLS = {
    "green": PatternFill("solid", fgColor="D1FADF"),
    "blue": PatternFill("solid", fgColor="DBEAFE"),
    "amber": PatternFill("solid", fgColor="FEF3C7"),
    "gray": PatternFill("solid", fgColor="F1F5F9"),
    "orange": PatternFill("solid", fgColor="FFE4D6"),
    "navy": PatternFill("solid", fgColor="C7D2FE"),
}
HEADER_ROW = 3
FIRST = 4
LAST = 200

wb = Workbook()


def sheet(title, subtitle, headers, widths):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = SUB
    for i, (h, w) in enumerate(zip(headers, widths, strict=True), start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    ws.row_dimensions[HEADER_ROW].height = 32
    return ws


def put_rows(ws, rows, *, numfmt=None, center_cols=()):
    for r, row in enumerate(rows, start=FIRST):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY, BORDER
            cell.alignment = CENTER if c in center_cols else WRAP
            if numfmt and c in numfmt:
                cell.number_format = numfmt[c]
    last = FIRST + len(rows) - 1
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(rows[0]))}{last}"
    return last


def status_rules(ws, col, last):
    rng = f"{col}{FIRST}:{col}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{FIRST},7)="Selesai"'], fill=FILLS["green"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{FIRST},8)="Snapshot"'], fill=FILLS["blue"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{FIRST},8)="Terbukti"'], fill=FILLS["green"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{FIRST},4)="Test"'], fill=FILLS["blue"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{FIRST},8)="Sebagian"'], fill=FILLS["amber"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{FIRST},5)="Belum"'], fill=FILLS["gray"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{FIRST},5)="Aktif"'], fill=FILLS["orange"]))


# ------------------------------------------------------------------ Agent Roster
roster_headers = ["ID", "Peran", "Jenis", "Model / tier rekomendasi", "Tanggung jawab", "Path milik (kepemilikan file)",
                  "Izin & batasan", "Input", "Output", "Direview oleh", "Wave aktif", "Status",
                  "Aktual token (juta)", "Aktual durasi (menit)", "Tool calls"]
ws_r = sheet("Agent Roster",
             "Siapa mengerjakan apa. Aktual = pemakaian nyata subagent pada Wave 1–2 (dari laporan harness). Kosong = belum berjalan / tidak terukur.",
             roster_headers, [8, 24, 16, 16, 46, 52, 28, 26, 26, 16, 12, 18, 12, 12, 10])
last_r = put_rows(ws_r, D.ROSTER, numfmt={13: "0.000", 14: "0.0", 15: "0"}, center_cols=(1, 11, 13, 14, 15))
status_rules(ws_r, "L", last_r)
tr = last_r + 1
ws_r.cell(row=tr, column=1, value="Total").font = BOLD
for col in (13, 14, 15):
    L = get_column_letter(col)
    c = ws_r.cell(row=tr, column=col, value=f"=SUM({L}{FIRST}:{L}{last_r})")
    c.font, c.border, c.alignment = BOLD, BORDER, CENTER
    c.number_format = {13: "0.000", 14: "0.0", 15: "0"}[col]
ws_r.cell(row=tr + 1, column=1,
          value="Catatan: token aktual AG-03 belum termasuk sesi hardening lanjutan (tidak dilaporkan). AG-00 (sesi utama) tidak diukur per WP.").font = SUB

# ------------------------------------------------------------------ WBS
wbs_headers = ["WP ID", "Phase", "Wave", "Modul", "Work package", "Deliverable / scope", "Owner (agent)",
               "Reviewer / gate", "Dependensi", "FR", "AC", "Prioritas", "Status", "Bukti / evidence",
               "Est. durasi agent (jam)", "Est. token (juta)", "Aktual token (juta)", "Aktual durasi (menit)", "Catatan"]
ws_w = sheet("WBS",
             "Work breakdown agentik per phase PRD §13. Semua kolom estimasi = [Menebak] (untuk WP selesai tanpa aktual = perkiraan retrospektif), dikalibrasi dari aktual Wave 1–2 (~0,33–0,73 juta token & 9–76 menit per agent-run).",
             wbs_headers, [11, 20, 7, 16, 30, 50, 16, 16, 18, 20, 18, 9, 18, 34, 11, 11, 11, 11, 28])
last_w = put_rows(ws_w, D.WBS, numfmt={15: "0.0", 16: "0.00", 17: "0.000", 18: "0.0"},
                  center_cols=(1, 3, 12, 15, 16, 17, 18))
status_rules(ws_w, "M", last_w)
dv = DataValidation(type="list", formula1='"Selesai,Snapshot (belum terintegrasi),Sebagian,Belum"', allow_blank=False)
ws_w.add_data_validation(dv)
dv.add(f"M{FIRST}:M{last_w}")
# blue font for hardcoded estimate inputs
for r in range(FIRST, last_w + 1):
    for col in (15, 16):
        ws_w.cell(row=r, column=col).font = Font(name=FONT, size=10, color="0000FF")
tw = last_w + 1
ws_w.cell(row=tw, column=1, value="Total").font = BOLD
for col in (15, 16, 17, 18):
    L = get_column_letter(col)
    c = ws_w.cell(row=tw, column=col, value=f"=SUM({L}{FIRST}:{L}{last_w})")
    c.font, c.border, c.alignment = BOLD, BORDER, CENTER
    c.number_format = {15: "0.0", 16: "0.00", 17: "0.000", 18: "0.0"}[col]
ws_w.cell(row=tw + 1, column=1,
          value="Kolom O–P (biru) = input estimasi yang dapat diubah. Kolom Q–R = aktual terukur. Status dapat dipilih dari daftar.").font = SUB

# ------------------------------------------------------------------ Timeline (formula-driven)
waves = [w[0] for w in D.WAVES]
tl_headers = ["Agent / peran", "Nama"] + waves + ["Jumlah WP"]
ws_t = sheet("Timeline",
             "Matriks wave × agent, dihitung otomatis dari sheet WBS (kolom Owner & Wave). ■ = agent memiliki ≥1 WP pada wave tsb.",
             tl_headers, [10, 26] + [8] * len(waves) + [11])
agents = [(r[0], r[1]) for r in D.ROSTER]
for i, (aid, name) in enumerate(agents):
    row = FIRST + i
    ws_t.cell(row=row, column=1, value=aid)
    ws_t.cell(row=row, column=2, value=name)
    for j, w in enumerate(waves):
        col = 3 + j
        L = get_column_letter(col)
        ws_t.cell(row=row, column=col,
                  value=f'=IF(COUNTIFS(WBS!$G${FIRST}:$G${LAST},"*"&$A{row}&"*",WBS!$C${FIRST}:$C${LAST},{L}${HEADER_ROW})>0,"■","")')
    ncol = 3 + len(waves)
    ws_t.cell(row=row, column=ncol, value=f'=COUNTIF(WBS!$G${FIRST}:$G${LAST},"*"&$A{row}&"*")')
    for c in range(1, ncol + 1):
        cell = ws_t.cell(row=row, column=c)
        cell.font, cell.border = BODY, BORDER
        cell.alignment = CENTER if c != 2 else WRAP
last_t = FIRST + len(agents) - 1
grid = f"C{FIRST}:{get_column_letter(2 + len(waves))}{last_t}"
ws_t.conditional_formatting.add(grid, CellIsRule(operator="equal", formula=['"■"'],
                                                 fill=PatternFill("solid", fgColor="FDBA9A"),
                                                 font=Font(name=FONT, color=ORANGE, bold=True)))
# wave legend table
lr = last_t + 3
legend_cols = [1, 2, 4, 6]
for c, h in zip(legend_cols, ["Wave", "Fokus", "Status", "Agent"], strict=True):
    cell = ws_t.cell(row=lr, column=c, value=h)
    cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
for i, (w, focus, st, ag) in enumerate(D.WAVES, start=1):
    for c, v in zip(legend_cols, [w, focus, st, ag], strict=True):
        cell = ws_t.cell(row=lr + i, column=c, value=v)
        cell.font, cell.border, cell.alignment = BODY, BORDER, Alignment(vertical="top")
ws_t.cell(row=lr - 1, column=1, value="Legenda wave").font = BOLD

# ------------------------------------------------------------------ RACI
ws_c = sheet("RACI",
             "R = mengerjakan, A = akuntabel/menyetujui, C = dikonsultasi, I = diberi tahu. Satu A per aktivitas.",
             ["Aktivitas"] + D.RACI_COLS, [38] + [7] * len(D.RACI_COLS))
rows = [[act] + [m.get(col, "") for col in D.RACI_COLS] for act, m in D.RACI]
last_c = put_rows(ws_c, rows, center_cols=tuple(range(2, len(D.RACI_COLS) + 2)))
rng = f"B{FIRST}:{get_column_letter(len(D.RACI_COLS) + 1)}{last_c}"
for letter, fill in (("R", "orange"), ("A", "navy"), ("C", "amber"), ("I", "gray")):
    ws_c.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{letter}"'], fill=FILLS[fill]))
ws_c.cell(row=last_c + 1, column=1, value="Cek: jumlah 'A' per baris").font = SUB
chk_col = len(D.RACI_COLS) + 2
ws_c.cell(row=HEADER_ROW, column=chk_col, value="Jumlah A").font = HDR_FONT
ws_c.cell(row=HEADER_ROW, column=chk_col).fill = HDR_FILL
ws_c.column_dimensions[get_column_letter(chk_col)].width = 10
for r in range(FIRST, last_c + 1):
    c = ws_c.cell(row=r, column=chk_col, value=f'=COUNTIF(B{r}:{get_column_letter(chk_col - 1)}{r},"A")')
    c.font, c.border, c.alignment = BODY, BORDER, CENTER

# ------------------------------------------------------------------ Traceability
ws_a = sheet("Traceability AC",
             "AC-01..AC-20 (PRD §12.1) → FR (FR_CATALOG, turunan) → WP pembangun → WP verifikasi E2E. 'Test ada' ≠ AC terpenuhi (bukti E2E/staging belum).",
             ["AC", "Kriteria", "FR terkait", "WP pembangun", "WP verifikasi", "Status"], [8, 44, 46, 22, 24, 22])
last_a = put_rows(ws_a, D.AC, center_cols=(1,))
status_rules(ws_a, "F", last_a)

# ------------------------------------------------------------------ Modules
ws_m = sheet("Status Modul", "20 modul PRD §5 → owner agent dan status kode snapshot (integrasi E2E: semua belum).",
             ["Modul", "Nama", "FR", "Prioritas", "Owner agent", "Status kode", "Gap utama"], [8, 32, 14, 10, 22, 20, 36])
last_m = put_rows(ws_m, D.MODULES, center_cols=(1, 4))
status_rules(ws_m, "F", last_m)

# ------------------------------------------------------------------ Gates
ws_g = sheet("Gates", "Gate persetujuan manusia. Agent tidak dapat melewati gate; User Proxy hanya merekomendasikan.",
             ["Gate", "Titik", "Penyetuju", "Kriteria lulus", "Bukti yang diminta", "Status"], [8, 30, 18, 52, 30, 22])
last_g = put_rows(ws_g, D.GATES, center_cols=(1,))
status_rules(ws_g, "F", last_g)

# ------------------------------------------------------------------ Risks
ws_k = sheet("Risiko", "Risiko PRD §15 (R-xx) + risiko baru dari build Wave 1–2 (RN-xx).",
             ["ID", "Risiko", "Severity", "Mitigasi", "Pemilik", "Status"], [8, 44, 12, 50, 16, 30])
last_k = put_rows(ws_k, D.RISKS, center_cols=(1, 3))
ws_k.conditional_formatting.add(f"C{FIRST}:C{last_k}", CellIsRule(operator="equal", formula=['"Critical"'],
                                                                   fill=PatternFill("solid", fgColor="FECACA")))
ws_k.conditional_formatting.add(f"C{FIRST}:C{last_k}", CellIsRule(operator="equal", formula=['"High"'],
                                                                   fill=FILLS["orange"]))

# ------------------------------------------------------------------ Ringkasan (first sheet)
ws = wb.active
ws.title = "Ringkasan"
ws.column_dimensions["A"].width = 46
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 60
ws["A1"] = "KodeKantor — WBS Pembagian Kerja Full Agentic"
ws["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
ws["A2"] = "Sumber: PRD KodeKantor v1.0 (21 Okt 2026) · Rancangan: docs/ARCHITECTURE.md & DECISIONS.md · Snapshot 21 Okt 2026"
ws["A2"].font = SUB
ws["A4"] = "Status program"
ws["A4"].font = BOLD
ws["A5"] = ("Rancangan (Phase 0) selesai disusun & direview independen; penguncian formal menunggu gate G0 (keputusan PO). "
            "Implementasi Wave 1–2 sempat berjalan lalu DIHENTIKAN atas instruksi Product Owner; kode tersimpan sebagai snapshot "
            "yang belum diintegrasikan end-to-end. Tidak ada klaim MVP selesai.")
ws.merge_cells("A5:C5")
ws["A5"].alignment = WRAP
ws["A5"].font = BODY
ws.row_dimensions[5].height = 44

metrics = [
    ("Jumlah work package", f"=COUNTA(WBS!A{FIRST}:A{last_w})", "0", "Semua phase 0–16"),
    ("WP Selesai", f'=COUNTIF(WBS!M{FIRST}:M{last_w},"Selesai")', "0", "Terverifikasi"),
    ("WP Snapshot (belum terintegrasi)", f'=COUNTIF(WBS!M{FIRST}:M{last_w},"Snapshot*")', "0", "Kode + test unit/integrasi lulus"),
    ("WP Sebagian", f'=COUNTIF(WBS!M{FIRST}:M{last_w},"Sebagian")', "0", "Ada gap tercatat"),
    ("WP Belum", f'=COUNTIF(WBS!M{FIRST}:M{last_w},"Belum")', "0", "Termasuk gate manusia"),
    ("Progres (Selesai+Snapshot) / total", f"=IFERROR((B8+B9)/B7,0)", "0.0%", "Bukan % fitur jadi; ukuran WP"),
    ("Agent AI (aktual + rencana)", f'=COUNTIF(\'Agent Roster\'!C{FIRST}:C{last_r},"AI agent*")', "0", "AG-00..AG-21"),
    ("Peran manusia", f'=COUNTIF(\'Agent Roster\'!C{FIRST}:C{last_r},"Manusia")', "0", "Product Owner, Operator"),
    ("Token aktual subagent terukur (juta)", f"=SUM('Agent Roster'!M{FIRST}:M{last_r})", "0.000", "Dari laporan harness"),
    ("Durasi aktual subagent terukur (menit, kumulatif)", f"=SUM('Agent Roster'!N{FIRST}:N{last_r})", "0.0", "Paralel: wall-clock jauh lebih pendek"),
    ("Estimasi token WP 'Belum' (juta) [Menebak]", f'=SUMIFS(WBS!P{FIRST}:P{last_w},WBS!M{FIRST}:M{last_w},"Belum")', "0.00", "Kalibrasi dari aktual"),
    ("Estimasi jam agent WP 'Belum' [Menebak]", f'=SUMIFS(WBS!O{FIRST}:O{last_w},WBS!M{FIRST}:M{last_w},"Belum")', "0.0", "Termasuk aktivitas manusia"),
    ("AC dengan test di snapshot", f'=COUNTIF(\'Traceability AC\'!F{FIRST}:F{last_a},"Test ada*")+COUNTIF(\'Traceability AC\'!F{FIRST}:F{last_a},"Terbukti*")', "0", "Dari 20 AC"),
    ("AC sebagian", f'=COUNTIF(\'Traceability AC\'!F{FIRST}:F{last_a},"Sebagian*")', "0", "AC-15, AC-19, AC-20"),
]
ws["A6"] = "Metrik"
ws["B6"] = "Nilai"
ws["C6"] = "Keterangan"
for c in ("A6", "B6", "C6"):
    ws[c].font, ws[c].fill, ws[c].border = HDR_FONT, HDR_FILL, BORDER
for i, (label, formula, fmt, note) in enumerate(metrics, start=7):
    ws.cell(row=i, column=1, value=label).font = BODY
    v = ws.cell(row=i, column=2, value=formula)
    v.number_format, v.font, v.alignment = fmt, BOLD, Alignment(horizontal="right")
    ws.cell(row=i, column=3, value=note).font = BODY
    for c in range(1, 4):
        ws.cell(row=i, column=c).border = BORDER
r0 = 7 + len(metrics) + 1
notes = [
    ("Cara pakai", ""),
    ("Agent Roster", "Peran, kepemilikan file (anti konflik paralel), izin, input/output, reviewer, pemakaian aktual."),
    ("WBS", "Work package per phase/wave dengan dependensi, FR/AC, status & bukti. Ubah status lewat daftar pilihan."),
    ("Timeline", "Matriks wave × agent otomatis dari WBS."),
    ("RACI", "Akuntabilitas; kolom 'Jumlah A' harus 1 per aktivitas."),
    ("Traceability AC", "AC → FR → WP pembangun → WP verifikasi."),
    ("Status Modul / Gates / Risiko", "Ringkasan modul M01–M20, gate persetujuan manusia, risiko & mitigasi."),
    ("Aturan kerja agent", "Path milik tegas; tanpa git write (hanya AG-00 commit); kontrak berubah → ADR; tidak boleh klaim test lulus tanpa output; gate manusia tidak dapat dilewati."),
    ("Label keyakinan", "Angka aktual = [Pasti] (diukur). Estimasi WP 'Belum' = [Menebak]."),
]
for i, (k, v) in enumerate(notes):
    a = ws.cell(row=r0 + i, column=1, value=k)
    a.font = BOLD
    a.alignment = Alignment(vertical="top")
    b = ws.cell(row=r0 + i, column=2, value=v)
    ws.merge_cells(start_row=r0 + i, start_column=2, end_row=r0 + i, end_column=3)
    b.alignment, b.font = WRAP, BODY
    ws.row_dimensions[r0 + i].height = 30 if v else 18

for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
    w.sheet_properties.tabColor = ORANGE if w.title in ("Ringkasan", "WBS") else NAVY
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
wb.save(OUT)
print("saved", OUT)
